from __future__ import annotations

import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from upress_workbench.citations import build_inventory
from upress_workbench.ingest import extract_chapter
from upress_workbench.packets import build_chapter_packet
from upress_workbench.review import author_items, load_skill_policy, run_rci, run_sei
from upress_workbench.schema_guard import SchemaViolation, validate
from upress_workbench.security import (
    SafetyError,
    SourceFile,
    scan_sources,
    sha256_file,
    validate_docx_archive,
    validate_output_parent,
)
from upress_workbench.state import StateStore


class StubClient:
    model = "glm-5.3-background"

    def complete_json(self, system_prompt, user_prompt, schema_name, schema):
        self.system_prompt = system_prompt
        self.user_prompt = user_prompt
        if "rci" in schema_name:
            return {
                "review_summary": "The year in the citation and reference entry does not match.",
                "release_status": "not_ready",
                "findings": [{
                    "skill_id": "RCI-01",
                    "rule_id": "RCI-REC-001",
                    "severity": "high",
                    "status": "needs_review",
                    "action": "BLOCK",
                    "confidence": 0.97,
                    "locator": "body.p0002",
                    "context_type": "body",
                    "exact_text": "(Smith, 2020)",
                    "reason": "The in-text year is 2020 and the reference-list year is 2021.",
                    "evidence": [{"source_id": "MANUSCRIPT-CH001", "locator": "body.p0002 and body.p0004", "relevance": "Direct comparison of preserved manuscript text", "access_status": "available"}],
                    "proposed_change": None,
                    "author_query": "Please confirm whether the citation or reference year is correct.",
                    "dependencies": [],
                }],
                "author_notes": [],
            }, {"model": self.model, "prompt_tokens": 100, "completion_tokens": 80, "total_tokens": 180}
        return {
            "review_summary": "One claim needs a source-scope check.",
            "release_status": "not_ready",
            "findings": [{
                "skill_id": "CEI-01",
                "rule_id": "SEI-CLAIM-001",
                "status": "source_unavailable",
                "risk_level": "moderate",
                "exact_text": "The intervention proved the claim.",
                "locator": "body.p0002",
                "context_type": "claim",
                "claim_or_object": "claim",
                "source_or_record": "R001",
                "source_access_status": "inaccessible",
                "comparison_or_basis": "The source text was not supplied.",
                "authority": "SEI-CLAIM-001",
                "reason": "A source comparison is required; this is not a plagiarism determination.",
                "owner": "Volume editor",
                "next_action": "Ask the author for the cited source passage and confirm scope.",
                "confidence": "high",
                "escalation": "editor",
                "closure_condition": "The editor compares the source passage and records a decision.",
                "release_effect": "hold",
                "proposed_change": None,
                "author_query": "Please supply the source passage supporting this statement.",
                "dependencies": ["RCI-CH001-001"],
            }],
            "author_notes": [],
        }, {"model": self.model, "prompt_tokens": 110, "completion_tokens": 90, "total_tokens": 200}


def make_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Synthetic Chapter for Review", level=0)
    document.add_paragraph("The intervention proved the claim (Smith, 2020).")
    document.add_heading("References", level=1)
    document.add_paragraph("Smith, J. (2021). A synthetic source. Example Journal, 1(1), 1-2.")
    document.save(path)


class WorkbenchCoreTests(unittest.TestCase):
    def test_schema_guard_rejects_extra_fields_without_coercion(self) -> None:
        schema = {"type": "object", "additionalProperties": False, "required": ["status"], "properties": {"status": {"type": "string", "enum": ["ready"]}}}
        with self.assertRaises(SchemaViolation):
            validate({"status": "ready", "helpful_extra": True}, schema)
        with self.assertRaises(SchemaViolation):
            validate({"status": 1}, schema)

    def test_skill_packages_match_lock(self) -> None:
        rci = load_skill_policy("reference-citation-integrity")
        sei = load_skill_policy("scholarly-editorial-integrity")
        self.assertEqual(rci.commit, "a49e9f686898030684134ba7a7f84ac8b68fc4dc")
        self.assertEqual(sei.commit, "098080247ab21173bbbe3da3ab90aa3cf860e38c")
        self.assertIn("RCI-REC-001", rci.rule_ids)
        self.assertIn("SEI-CLAIM-001", sei.rule_ids)

    def test_output_inside_source_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "chapters"
            output = source / "reviews"
            source.mkdir()
            output.mkdir()
            with self.assertRaises(SafetyError):
                validate_output_parent(output, source.resolve())

    def test_recovery_only_interrupts_orphaned_processes(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            store = StateStore(root / "state.sqlite3")
            record = {
                "run_id": "RUN-1", "created_at": "2026-09-09T00:00:00Z", "updated_at": "2026-09-09T00:00:00Z",
                "status": "running", "input_root": str(root), "output_root": str(root / "output"), "project_label": "QA",
                "case_owner": "Editor", "classification": "public", "provider_approved": 1, "approval_at": "2026-09-09T00:00:00Z",
                "model": "glm-5.3-background", "crossref_enabled": 0, "concurrency": 1, "run_mode": "baseline",
                "worker_pid": os.getpid(), "heartbeat_at": "2026-09-09T00:00:00Z", "manifest_hash": "a", "plan_hash": "b",
                "error": None, "budget_json": "{}",
            }
            store.create_run(record, [])
            self.assertEqual(store.recover_orphaned_runs(), 0)
            store.update_run("RUN-1", worker_pid=99999999)
            self.assertEqual(store.recover_orphaned_runs(), 1)
            self.assertEqual(store.get_run("RUN-1")["status"], "interrupted")

    def test_docx_with_doctype_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "hostile.docx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("[Content_Types].xml", "<Types/>")
                archive.writestr("word/document.xml", "<!DOCTYPE x [<!ENTITY y 'bad'>]><document>&y;</document>")
            with self.assertRaises(SafetyError):
                validate_docx_archive(path)

    def test_docx_with_external_template_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "hostile.docx"
            relationships = """<?xml version='1.0' encoding='UTF-8'?>
<Relationships xmlns='http://schemas.openxmlformats.org/package/2006/relationships'>
  <Relationship Id='rId1' Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate' Target='file:///C:/hostile.dotm' TargetMode='External'/>
</Relationships>"""
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("[Content_Types].xml", "<Types/>")
                archive.writestr("word/document.xml", "<document/>")
                archive.writestr("word/_rels/document.xml.rels", relationships)
            with self.assertRaises(SafetyError):
                validate_docx_archive(path)

    def test_scan_and_inventory_preserve_source(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "chapter.docx"
            make_docx(path)
            original_hash = sha256_file(path)
            sources = scan_sources(root)
            self.assertEqual(len(sources), 1)
            chapter = extract_chapter(sources[0], "CH001")
            inventory = build_inventory(chapter)
            self.assertEqual(len(inventory["references"]), 1)
            self.assertGreaterEqual(len(inventory["citations"]), 1)
            self.assertIn("C001", inventory["unmatched_citation_ids"])
            self.assertEqual(sha256_file(path), original_hash)

    def test_offline_review_and_author_packet(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source_path = root / "chapter.docx"
            make_docx(source_path)
            source_hash = sha256_file(source_path)
            source = SourceFile(source_path, "chapter.docx", source_path.stat().st_size, source_hash)
            chapter = extract_chapter(source, "CH001")
            inventory = build_inventory(chapter)
            client = StubClient()
            rci_policy = load_skill_policy("reference-citation-integrity")
            sei_policy = load_skill_policy("scholarly-editorial-integrity")
            rci, rci_work = run_rci(client, rci_policy, chapter, inventory, [], "UniversityPress-APA7-pilot-v0.1", "baseline")
            sei, sei_work = run_sei(client, sei_policy, chapter, inventory, rci, [], "Volume editor", "unpublished_approved")
            items = author_items(rci, rci_work, sei, sei_work)
            output_root = root / "output"
            output_root.mkdir()
            receipt = build_chapter_packet(chapter, inventory, items, rci, rci_work, sei, sei_work, output_root)
            reviewed = Path(receipt["author_packet"]["path"])
            self.assertTrue(reviewed.is_file())
            self.assertTrue(receipt["source_unchanged"])
            self.assertEqual(sha256_file(source_path), source_hash)
            self.assertGreaterEqual(receipt["author_packet"]["comments_structurally_present"], 1)
            self.assertNotIn("plagiarism", json.dumps(sei, ensure_ascii=False).casefold())
            with zipfile.ZipFile(reviewed) as archive:
                self.assertIn("word/comments.xml", archive.namelist())
                document_xml = archive.read("word/document.xml")
                self.assertIn(b"commentRangeStart", document_xml)
                self.assertIn(b"University Press Review Notes", document_xml)

    def test_packet_build_blocks_source_drift(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source_path = root / "chapter.docx"
            make_docx(source_path)
            source = SourceFile(source_path, "chapter.docx", source_path.stat().st_size, sha256_file(source_path))
            chapter = extract_chapter(source, "CH001")
            inventory = build_inventory(chapter)
            client = StubClient()
            rci = run_rci(client, load_skill_policy("reference-citation-integrity"), chapter, inventory, [], "UniversityPress-APA7-pilot-v0.1", "baseline")
            sei = run_sei(client, load_skill_policy("scholarly-editorial-integrity"), chapter, inventory, rci[0], [], "Volume editor", "unpublished_approved")
            Document().save(source_path)
            output_root = root / "output"
            output_root.mkdir()
            with self.assertRaises(SafetyError):
                build_chapter_packet(chapter, inventory, author_items(rci[0], rci[1], sei[0], sei[1]), rci[0], rci[1], sei[0], sei[1], output_root)


if __name__ == "__main__":
    unittest.main()
