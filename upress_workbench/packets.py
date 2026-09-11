from __future__ import annotations

import html
import os
import re
import shutil
import uuid
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .artifacts import atomic_write_bytes, atomic_write_json, object_sha256, utc_now
from .ingest import Chapter
from .security import SafetyError, read_approved_bytes, safe_slug, sha256_file


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    tr_pr.append(no_split)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def _comment_text(item: dict[str, Any]) -> str:
    parts = [f"{item['id']} | {item['family']} | {item['status']}", item["reason"]]
    if item.get("proposed_change"):
        parts.append("Proposed change: " + item["proposed_change"])
    if item.get("author_query"):
        parts.append("Author query: " + item["author_query"])
    parts.append("This is a proposal for human editorial review, not an accepted change.")
    return "\n\n".join(parts)[:4000]


def _append_docx_notes(document: Document, items: list[dict[str, Any]], review_summary: str) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    heading = document.add_paragraph("University Press Review Notes", style="Heading 1")
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    intro = document.add_paragraph()
    intro.add_run("Purpose. ").bold = True
    intro.add_run(
        "These notes identify proposed citation and scholarly-integrity revisions for author and editor review. "
        "The source chapter was preserved. No proposed change has been accepted and this packet does not authorize publication."
    )
    if review_summary:
        summary = document.add_paragraph()
        summary.add_run("Review summary. ").bold = True
        summary.add_run(review_summary)
    if not items:
        document.add_paragraph("No author-facing changes or queries were proposed in this review pass.")
        return

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.0), Inches(1.25), Inches(3.45), Inches(4.15)]
    headers = ["ID", "Location", "Request", "Basis"]
    for index, (cell, label) in enumerate(zip(table.rows[0].cells, headers, strict=True)):
        cell.width = widths[index]
        cell.text = label
        _set_cell_shading(cell, "1F3A5F")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    _set_repeat_header(table.rows[0])
    _keep_row_together(table.rows[0])
    for row_index, item in enumerate(items):
        row = table.add_row()
        _keep_row_together(row)
        cells = row.cells
        request = item.get("proposed_change") or item.get("author_query") or "Review with the editor."
        values = [item["id"], item["locator"], request, item["reason"]]
        for cell_index, (cell, value) in enumerate(zip(cells, values, strict=True)):
            cell.width = widths[cell_index]
            cell.text = str(value)
            if row_index % 2:
                _set_cell_shading(cell, "F3F6FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def _annotate_docx(chapter: Chapter, items: list[dict[str, Any]], output: Path, review_summary: str) -> dict[str, Any]:
    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_bytes(output, read_approved_bytes(Path(chapter.source_path), chapter.sha256))
    document = Document(output)
    by_locator = {block.locator: block for block in chapter.blocks if block.paragraph_index is not None}
    comments_added = 0
    comment_errors = []
    for item in items[:50]:
        block = by_locator.get(item["locator"])
        if block is None or block.paragraph_index is None or block.paragraph_index >= len(document.paragraphs):
            continue
        paragraph = document.paragraphs[block.paragraph_index]
        if not paragraph.runs:
            continue
        try:
            document.add_comment(paragraph.runs, text=_comment_text(item), author="University Press Review", initials="UP")
            comments_added += 1
        except Exception as exc:  # appendix below remains the visible fallback
            comment_errors.append(type(exc).__name__)
    _append_docx_notes(document, items, review_summary)
    document.save(output)
    structural_comments = 0
    with zipfile.ZipFile(output) as archive:
        if "word/comments.xml" in archive.namelist():
            payload = archive.read("word/comments.xml")
            structural_comments = len(re.findall(br"<w:comment\b", payload))
    return {
        "path": str(output),
        "sha256": sha256_file(output),
        "comments_requested": min(50, len(items)),
        "comments_added": comments_added,
        "comments_structurally_present": structural_comments,
        "comment_fallback": comments_added < min(50, len(items)),
        "comment_error_types": sorted(set(comment_errors)),
        "appendix_items": len(items),
    }


def _annotate_text(chapter: Chapter, items: list[dict[str, Any]], output: Path, review_summary: str) -> dict[str, Any]:
    try:
        source = read_approved_bytes(Path(chapter.source_path), chapter.sha256).decode("utf-8", errors="strict")
    except UnicodeDecodeError as exc:
        raise SafetyError("Text and Markdown chapters must be UTF-8 encoded.") from exc
    lines = [source.rstrip(), "", "---", "", "# University Press Review Notes", ""]
    lines.append("These are proposed changes and questions for author and editor review. The source chapter was preserved. No change has been accepted and this packet does not authorize publication.")
    if review_summary:
        lines.extend(["", "## Review summary", "", review_summary])
    if not items:
        lines.extend(["", "No author-facing changes or queries were proposed in this review pass."])
    for item in items:
        lines.extend(["", f"## {item['id']} at {item['locator']}", "", f"Status: {item['status']}", "", item["reason"]])
        if item.get("proposed_change"):
            lines.extend(["", "Proposed change:", "", item["proposed_change"]])
        if item.get("author_query"):
            lines.extend(["", "Author query:", "", item["author_query"]])
    atomic_write_bytes(output, ("\n".join(lines).rstrip() + "\n").encode("utf-8"))
    return {"path": str(output), "sha256": sha256_file(output), "comments_requested": 0, "comments_added": 0, "comments_structurally_present": 0, "comment_fallback": True, "comment_error_types": [], "appendix_items": len(items)}


def _author_queries(items: list[dict[str, Any]], chapter: Chapter) -> str:
    lines = [f"# Author Queries for {Path(chapter.relative_path).stem}", "", "These queries require author or editor review. They are not accepted changes and do not authorize publication.", ""]
    if not items:
        lines.append("No author-facing queries were proposed in this pass.")
    for item in items:
        lines.extend([f"## {item['id']}", "", f"Location: `{item['locator']}`", "", item.get("author_query") or item.get("proposed_change") or "Please review this item with the editor.", "", f"Basis: {item['reason']}", ""])
    return "\n".join(lines).rstrip() + "\n"


def _review_html(chapter: Chapter, items: list[dict[str, Any]], rci: dict[str, Any], sei: dict[str, Any]) -> str:
    cards = []
    for item in items:
        request = item.get("proposed_change") or item.get("author_query") or "Review with the editor."
        cards.append(
            f"<article><div class='eyebrow'>{html.escape(item['family'])} · {html.escape(item['severity'])} · {html.escape(item['status'])}</div>"
            f"<h2>{html.escape(item['id'])}</h2><p class='locator'>{html.escape(item['locator'])}</p>"
            f"<h3>Requested response</h3><p>{html.escape(request)}</p><h3>Basis</h3><p>{html.escape(item['reason'])}</p></article>"
        )
    if not cards:
        cards.append("<article><h2>No author-facing items</h2><p>This review pass produced no proposed changes or author queries.</p></article>")
    title = html.escape(Path(chapter.relative_path).stem)
    return f"""<!doctype html><html lang='en'><head><meta charset='utf-8'><meta http-equiv='Content-Security-Policy' content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'"><meta name='viewport' content='width=device-width,initial-scale=1'><title>{title} review</title><style>
body{{margin:0;background:#f4f1ea;color:#162335;font:16px/1.6 Georgia,serif}}main{{max-width:980px;margin:auto;padding:56px 24px}}header{{border-bottom:1px solid #b9afa1;padding-bottom:28px;margin-bottom:28px}}h1,h2,h3{{font-family:Segoe UI,Arial,sans-serif;color:#132d4f}}h1{{font-size:2.4rem;line-height:1.1}}h2{{margin:.25rem 0}}h3{{font-size:.82rem;text-transform:uppercase;letter-spacing:.09em;margin-top:1.3rem}}.status{{display:flex;gap:10px;flex-wrap:wrap}}.pill{{background:#fff;border:1px solid #c9c2b8;border-radius:999px;padding:5px 11px;font:600 13px Segoe UI,Arial,sans-serif}}article{{background:white;border-left:5px solid #b5683a;box-shadow:0 6px 24px #23364d16;margin:18px 0;padding:24px 28px}}.eyebrow,.locator{{font:600 12px Segoe UI,Arial,sans-serif;text-transform:uppercase;letter-spacing:.08em;color:#667386}}.notice{{font-style:italic;color:#4c596a}}
</style></head><body><main><header><div class='eyebrow'>University Press author review packet</div><h1>{title}</h1><p class='notice'>Proposals and queries only. The source chapter is unchanged. Human editorial review is required before any change or release.</p><div class='status'><span class='pill'>RCI {html.escape(rci['release_status'])}</span><span class='pill'>SEI {html.escape(sei['release_status'])}</span><span class='pill'>{len(items)} author-facing items</span></div></header>{''.join(cards)}</main></body></html>"""


def build_chapter_packet(
    chapter: Chapter,
    inventory: dict[str, Any],
    items: list[dict[str, Any]],
    rci_result: dict[str, Any],
    rci_workbench: dict[str, Any],
    sei_result: dict[str, Any],
    sei_workbench: dict[str, Any],
    output_root: Path,
) -> dict[str, Any]:
    folder_name = f"{chapter.chapter_id}-{safe_slug(Path(chapter.relative_path).stem)}"
    chapters_root = output_root / "chapters"
    chapters_root.mkdir(parents=True, exist_ok=True)
    root = chapters_root / folder_name
    if root.exists():
        raise SafetyError(f"A completed or legacy partial packet already exists for {chapter.chapter_id}; preserve it and resolve the duplicate before retrying.")
    staging = chapters_root / f".{folder_name}.staging-{uuid.uuid4().hex[:8]}"
    author_root = staging / "author-review"
    editorial_root = staging / "editorial-record"
    author_root.mkdir(parents=True, exist_ok=False)
    editorial_root.mkdir(parents=True, exist_ok=False)

    try:
        review_summary = (
            f"Reference and Citation Integrity: {rci_result['release_status']} with {len(rci_result['findings'])} findings. "
            f"Scholarly and Editorial Integrity: {sei_result['release_status']} with {len(sei_result['findings'])} findings. "
            "Review each proposal or query below. Complete evidence, uncertainty, and provenance records are retained in the editorial-record folder."
        )
        source = Path(chapter.source_path)
        if chapter.extension == ".docx":
            reviewed_name = f"{source.stem}__AUTHOR_REVIEW.docx"
            reviewed_path = author_root / reviewed_name
            annotated = _annotate_docx(chapter, items, reviewed_path, review_summary)
        else:
            reviewed_name = f"{source.stem}__AUTHOR_REVIEW.md"
            reviewed_path = author_root / reviewed_name
            annotated = _annotate_text(chapter, items, reviewed_path, review_summary)
        queries_name = f"{source.stem}__AUTHOR_QUERIES.md"
        queries_path = author_root / queries_name
        atomic_write_bytes(queries_path, _author_queries(items, chapter).encode("utf-8"))
        review_html_path = author_root / "OPEN_REVIEW.html"
        atomic_write_bytes(review_html_path, _review_html(chapter, items, rci_result, sei_result).encode("utf-8"))

        atomic_write_json(editorial_root / "citation-inventory.json", inventory)
        atomic_write_json(editorial_root / "rci-run-result.json", rci_result)
        atomic_write_json(editorial_root / "rci-workbench-record.json", rci_workbench)
        atomic_write_json(editorial_root / "sei-run-result.json", sei_result)
        atomic_write_json(editorial_root / "sei-workbench-record.json", sei_workbench)
        source_unchanged = sha256_file(Path(chapter.source_path)) == chapter.sha256
        if not source_unchanged:
            raise SafetyError("Source changed while the author packet was being assembled.")

        annotated["path"] = str(root / "author-review" / reviewed_name)
        receipt = {
            "schema_version": 1,
            "created_at": utc_now(),
            "chapter_id": chapter.chapter_id,
            "source": {"path": chapter.source_path, "relative_path": chapter.relative_path, "sha256": chapter.sha256, "size": chapter.size},
            "source_unchanged": source_unchanged,
            "author_packet": annotated,
            "author_queries": {"path": str(root / "author-review" / queries_name), "sha256": sha256_file(queries_path)},
            "review_html": {"path": str(root / "author-review" / "OPEN_REVIEW.html"), "sha256": sha256_file(review_html_path)},
            "finding_counts": {"rci": len(rci_result["findings"]), "sei": len(sei_result["findings"]), "author_facing": len(items)},
            "release_status": {"rci": rci_result["release_status"], "sei": sei_result["release_status"]},
            "publication_authorized": False,
        }
        receipt["receipt_sha256"] = object_sha256(receipt)
        atomic_write_json(editorial_root / "chapter-receipt.json", receipt)
        os.replace(staging, root)
        return receipt
    except Exception:
        if staging.exists() and staging.parent == chapters_root:
            shutil.rmtree(staging, ignore_errors=True)
        raise


def run_summary_html(run: dict[str, Any], chapters: list[dict[str, Any]]) -> str:
    rows = []
    for chapter in chapters:
        packet = chapter.get("packet") or {}
        rel = packet.get("review_html_relative")
        link = f"<a href='{html.escape(rel)}'>Open author review</a>" if rel else "Not available"
        rows.append(f"<tr><td>{html.escape(chapter['relative_path'])}</td><td>{html.escape(chapter['status'])}</td><td>{chapter.get('rci_findings', 0)}</td><td>{chapter.get('sei_findings', 0)}</td><td>{link}</td></tr>")
    return f"""<!doctype html><html lang='en'><head><meta charset='utf-8'><meta http-equiv='Content-Security-Policy' content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'"><meta name='viewport' content='width=device-width,initial-scale=1'><title>University Press review run</title><style>
body{{margin:0;background:#f4f1ea;color:#162335;font:15px/1.5 Segoe UI,Arial,sans-serif}}main{{max-width:1100px;margin:auto;padding:50px 24px}}h1{{font:700 38px/1.15 Georgia,serif;color:#132d4f}}.notice{{font-family:Georgia,serif;font-style:italic}}table{{width:100%;border-collapse:collapse;background:white;box-shadow:0 6px 24px #23364d16}}th,td{{padding:13px 14px;border-bottom:1px solid #d7d3cc;text-align:left}}th{{background:#1f3a5f;color:white}}a{{color:#9b4d24;font-weight:700}}.meta{{display:flex;gap:10px;flex-wrap:wrap;margin:24px 0}}.pill{{background:white;border:1px solid #cbc4ba;border-radius:999px;padding:6px 12px}}
</style></head><body><main><p>UNIVERSITY PRESS WORKBENCH</p><h1>{html.escape(run.get('project_label') or 'Citation review run')}</h1><p class='notice'>Review proposals only. Original chapters were not overwritten. Publication remains a human decision.</p><div class='meta'><span class='pill'>Run {html.escape(run['run_id'])}</span><span class='pill'>Status {html.escape(run['status'])}</span><span class='pill'>{len(chapters)} chapters</span></div><table><thead><tr><th>Chapter</th><th>Status</th><th>RCI</th><th>SEI</th><th>Packet</th></tr></thead><tbody>{''.join(rows)}</tbody></table></main></body></html>"""
